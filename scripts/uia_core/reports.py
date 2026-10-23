from __future__ import annotations

import html
import json
import zipfile
from pathlib import Path
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .common import atomic_write, elapsed, file_hash, inside, now, read_json, summary
from .secrets import drop_unsafe_artifacts, scrub_text, unsafe_attachment
from . import metrics as run_metrics


STATUS_ZH = {"remote_pending": "待远程检查", "remote_cached": "公开文件仍在缓存期", "remote_unreachable": "远程不可达或无法验证","passed": "通过", "business_failure": "业务失败", "save_not_persisted": "保存未生效（后台读回不一致）", "automation_failure": "自动化失败",
             "environment_blocked": "环境阻断", "product_confirmation": "待业务确认", "not_suitable": "不适合自动化", "suspended": "挂起（流程/知识缺口）",
             "interrupted": "已中断", "stable": "已稳定", "unstable": "未稳定", "candidate": "待复验",
             "not_run": "未执行", "pending": "待处理", "runnable": "可执行", "running": "执行中",
             "waiting_human": "等待前置", "finished": "已完成", "open": "待回答", "answered": "已回答",
             "closed": "已关闭", "deferred": "待异步审核", "active": "进行中", "delivery_pending": "待交付",
             "completed": "已完成", "waiting": "等待前置"}


def zh(value):
    return STATUS_ZH.get(value, value) if value is not None else "未执行"


def opening(m, metrics, review):
    from collections import Counter
    counts = Counter(zh(c.get('result')) if c.get('work_status') != 'waiting_human' else '等待前置' for c in m['cases'].values())
    conclusion = '本次共 ' + str(len(m['cases'])) + ' 条：' + '，'.join(str(n) + ' 条' + label for label, n in counts.items()) + '。'
    waiting = [cid + '：' + c.get('classification', {}).get('reason', '见审核单') for cid, c in m['cases'].items() if c.get('work_status') == 'waiting_human']
    duplicates = [cid + '：' + c['duplicate_script_warning'] + ' 执行 ' + (c.get('registration') or {}).get('script', '')
                  for cid, c in m['cases'].items() if c.get('duplicate_script_warning')]
    if duplicates:
        waiting.append('重复脚本选择：' + '；'.join(duplicates))
    remote = [cid for cid, c in m['cases'].items() if c.get('work_status') == 'remote_pending']
    if remote:
        waiting.append('公开文件待远程验证：' + '、'.join(remote) + '；见 remote-checks/request.json，不计通过。')
    needs = waiting + ([f"异步审核单 {len(review['items'])} 项；断言建议默认不应用。"] if review['items'] else [])
    metrics_text = (f"清单问题 {metrics['checklist_questions']}（入口问题 {metrics.get('entry_question_count', 0)}），清单回答 {metrics['checklist_answers']}，中途提问 {metrics['mid_run_questions']}，"
                    f"计划外人工介入 {metrics['unplanned_interventions']}，会话轮次 {metrics['session_turns']}；"
                    f"脚本实际执行覆盖率 {metrics['script_coverage']['ratio']:.1%}；"
                    f"越权 {metrics['scope_violations']}，泄露 {metrics['leak_hits']}，未审核断言变更 {metrics['unreviewed_assertion_changes']}。")
    if metrics.get('re_explored_known_pages', 0):
        metrics_text += f"流程问题：已知页面重复探索 {metrics['re_explored_known_pages']} 次（re_explored_known_pages）。"
    if metrics.get('exploration_for_missing_controls', 0):
        metrics_text += f"缺控件探索 {metrics['exploration_for_missing_controls']} 次（exploration_for_missing_controls）。"
    if metrics.get('knowledge_injected_tokens', {}).get('count'):
        kb = metrics['knowledge_injected_tokens']
        metrics_text += f"知识注入 token p50/p95：{kb['p50']}/{kb['p95']}（{kb['count']} 条用例）。"
    writeback = m.get('knowledge_writeback', {})
    if writeback.get('count'):
        metrics_text += ' 本轮知识新增或更新：' + '、'.join(x['name'] for x in writeback.get('added_or_updated', [])) + '。'
    if writeback.get('partial_count'):
        metrics_text += f" 执行成功步骤的部分知识回写 {writeback['partial_count']} 项。"
    checklist = m.get('questions', {}).get('run-preconditions', {})
    if (checklist.get('entry_count', 0) and checklist.get('memory_answer_count') == checklist.get('entry_count')
            and checklist.get('status') == 'answered' and metrics['checklist_questions'] == 0):
        metrics_text += f" 清单全部来自环境记忆（{checklist['memory_answer_count']} 项），本次 0 问。"
    return [{'title': '结论', 'text': conclusion}, {'title': '需要你处理', 'text': '\n'.join(needs) or '无。'},
            {'title': '流程指标', 'text': metrics_text}]


def waiting_seconds(m):
    entries = [q for q in m["questions"].values()]
    entries += [prior for q in m["questions"].values() for prior in q.get("previous_answers", []) if prior.get("created_at")]
    intervals = sorted((q["created_at"], q.get("answered_at") or now()) for q in entries)
    merged = []
    for start, end in intervals:
        if merged and start <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(end, merged[-1][1]))
        else:
            merged.append((start, end))
    return sum(elapsed(start, end) for start, end in merged)


def brief_error(attempt, cfg=None):
    return scrub_text(" ".join(attempt.get("errors") or []), cfg)[:180]


def make_report(base, m):
    from .knowledge import writeback as knowledge_writeback
    knowledge_writeback(base, m)
    from .experience import review as build_review
    review = build_review(base, m)
    removed = drop_unsafe_artifacts(base / "batches")
    out = base / "reports"
    out.mkdir(parents=True, exist_ok=True)
    for c in m["cases"].values():
        for attempt in c.get("attempts") or []:
            attempt["errors"] = [scrub_text(item, m["config"]) for item in (attempt.get("errors") or [])]
            attempt["attachments"] = [item for item in (attempt.get("attachments") or []) if not unsafe_attachment(item)]
    records = []
    for c in m["cases"].values():
        records.append({k: c.get(k) for k in ("case_id", "original_id", "title", "module", "priority",
                       "sheet", "rows", "stage", "work_status", "result", "stability", "pass_streak",
                       "flow_ids", "fixture_ids", "repair_rounds", "attempts", "feedback", "classification", "case_spec", "spec_first_lint", "spec_lint", "remote_status")})
        records[-1]['result_label'] = zh(c.get('result'))
        records[-1]['script_binding'] = c.get('script_binding')
        records[-1]['read_back'] = (c.get('attempts') or [{}])[-1].get('read_back')
        records[-1]['duplicate_script_warning'] = c.get('duplicate_script_warning')
        records[-1]['work_status_label'] = zh(c.get('work_status'))
        records[-1]['stability_label'] = zh(c.get('stability'))
        records[-1]["script"] = (c.get("registration") or {}).get("script")
        attempt = c["attempts"][-1] if c.get("attempts") else None
        classification = c.get("classification") or {}
        records[-1]["reason"] = classification.get("reason", "") if isinstance(classification, dict) else ""
        if c.get("result") == "save_not_persisted":
            records[-1]["reason"] = "保存未生效（后台读回不一致）"
        elif attempt and attempt.get("executed") and c.get("result") == "passed":
            records[-1]["reason"] = brief_error(attempt, m["config"])
        elif attempt and attempt.get("executed") and c.get("result") in {"business_failure", "save_not_persisted", "automation_failure", "interrupted"}:
            records[-1]["reason"] = brief_error(attempt, m["config"]) or records[-1]["reason"]
        elif not records[-1]["reason"] and attempt:
            records[-1]["reason"] = brief_error(attempt, m["config"])
        records[-1]["reason"] = scrub_text(records[-1]["reason"], m["config"])
        evidence = records[-1].get('read_back')
        if evidence and evidence.get('ok'):
            records[-1]["reason"] = (records[-1]["reason"] + "；" if records[-1]["reason"] else "") + "后台读回：" + str(evidence.get("actual", "")) + " ✓"
        records[-1]["media"] = ['../' + a['redacted_screenshot'] for a in c.get('attempts', []) if a.get('redacted_screenshot')]
    m["final_report_at"] = now()
    m["leak_hits"] = run_metrics.leak_scan(base, m["config"])
    flow_metrics = run_metrics.calculate(base, m, m["final_report_at"])
    atomic_write(base / "metrics.json", flow_metrics)
    atomic_write(out / "metrics.json", flow_metrics)
    blocks = opening(m, flow_metrics, review)
    opening_md = "\n\n".join("## " + b['title'] + "\n\n" + b['text'] for b in blocks) + "\n"
    (out / 'summary.md').write_text(opening_md, encoding='utf-8')
    s = summary(m)
    metrics = {"elapsed_seconds": elapsed(m["created_at"]), "human_wait_seconds": waiting_seconds(m),
               "first_result_seconds": elapsed(m["created_at"], m["first_result_at"]) if m.get("first_result_at") else None,
               "first_pass_seconds": elapsed(m["created_at"], m["first_pass_at"]) if m.get("first_pass_at") else None,
               "exploration_seconds": sum(a.get("seconds", 0) for a in m["exploration"]),
               "exploration_attempts": len(m["exploration"]),
               "human_questions": len(m["questions"]),
               "cases_with_asset_reuse": sum(bool((c.get("registration") or {}).get("reused_assets")) for c in m["cases"].values()),
               "repair_rounds": sum(c["repair_rounds"] for c in m["cases"].values()),
               **flow_metrics}
    questions = [{k: q.get(k) for k in ("id", "kind", "subject", "prompt", "case_ids", "status", "attempted")}
                 for q in m["questions"].values()]
    payload = {"schema_version": "uia-results.v1", "run_id": m["run_id"], "generated_at": now(),
               "remote_checks": m.get("remote_checks", {}), "opening": blocks, "summary": s, "metrics": metrics, "flow_metrics": flow_metrics, "source": m["sources"], "cases": records,
               "questions": questions, "unresolved_sheets": [x for src in m["sources"] for x in src.get("unresolved_sheets", [])],
               "source_update_required": m.get("source_update_required", False), "delivery_complete": False}
    diff = {"schema_version": "uia-asset-review.v1", "writeback": "not_applied", "cases": []}
    diff['entry_hints'] = {'file': 'assets/playwright-runner/knowledge/entries.yaml',
                           'changes': m.get('entry_hint_changes', {})}
    for c in m["cases"].values():
        if c.get("feedback"):
            file = inside(base, c["feedback"])
            diff["cases"].append({"case_id": c["case_id"], "review": c["feedback"],
                                  "sha256": file_hash(file), "content": read_json(file)})
    atomic_write(out / "asset-review.json", diff)
    atomic_write(out / "results.json", payload)
    wb = Workbook()
    overview = wb.active
    overview.title = "报告开头"
    for block in blocks:
        overview.append([block['title'], block['text']])
    sheet = wb.create_sheet("逐条结果")
    sheet.title = "逐条结果"
    sheet.append(["用例ID", "原始ID", "工作表", "源行号", "标题", "优先级", "进度", "结果", "稳定性", "连续成功", "脚本", "执行次数", "最近证据", "原因"])
    for r in records:
        sheet.append([r["case_id"], r["original_id"], r["sheet"], ",".join(map(str, r["rows"])),
                      r["title"], r["priority"], zh(r["work_status"]), zh(r["result"]), zh(r["stability"]), r["pass_streak"],
                      r["script"], len(r["attempts"]), r["attempts"][-1]["evidence"] if r["attempts"] else "", r.get("reason", "")])
    stats = wb.create_sheet("汇总")
    stats.append(["指标", "值"])
    for key, value in {**s, **metrics}.items():
        stats.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    qs = wb.create_sheet("人工问题")
    qs.append(["问题ID", "类型", "问题", "影响用例", "状态", "已尝试"])
    for q in questions:
        qs.append([q["id"], q["kind"], q["prompt"], ", ".join(q["case_ids"]), zh(q["status"]), "\n".join(q["attempted"])])
    for ws in wb:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="23415A")
        for row in ws.iter_rows():
            for cell in row:
                # Source text must never become an executable spreadsheet formula.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(64, max(14, max(len(str(c.value or "")) for c in col[:100]) + 2))
    wb.save(out / "results.xlsx")
    escape = lambda value: html.escape(str(value if value is not None else ""))
    def link(ref):
        if not ref:
            return ""
        return '<a href="../' + quote(str(ref), safe="/") + '">' + escape(ref) + '</a>'
    rows = []
    for r in records:
        source = escape(r["sheet"]) + ":" + escape(",".join(map(str, r["rows"])))
        shots = " ".join('<a href="' + escape(ref) + '">截图</a>' for ref in r.get("media") or [])
        rows.append("<tr>" + "".join("<td>" + x + "</td>" for x in [escape(r["case_id"]), source, escape(r["title"]),
                    escape(zh(r["work_status"])), escape(zh(r["result"])), escape(zh(r["stability"])), str(r["pass_streak"]),
                    link(r["script"]), link(r["attempts"][-1]["evidence"] if r["attempts"] else None),
                    escape(r.get("reason", "")) + (" " + shots if shots else "")]) + "</tr>")
    body = """<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width">
<title>UI 自动化结果</title><style>body{font:15px system-ui;margin:32px;color:#193348;background:#f5f8fb}
table{border-collapse:collapse;width:100%;background:white}th,td{padding:10px;text-align:left;border:1px solid #dbe4eb;vertical-align:top}
th{background:#23415a;color:white}a{color:#0069a3}pre{white-space:pre-wrap}section{overflow:auto}input{padding:10px;min-width:260px;margin:12px 0}</style>
<h1>UI 自动化结果</h1>"""
    body += ''.join('<h2>' + escape(b['title']) + '</h2><p>' + escape(b['text']).replace('\n', '<br>') + '</p>' for b in blocks)
    passes = m["config"]["execution"].get("stability_passes", 2)
    standard = "每日交付：同版本一次成功即可。" if passes == 1 else "夜间回归：同一版本独立运行两次成功。"
    body += "<p>运行：" + escape(m["run_id"]) + "；源用例总数：" + str(s["total"]) + "。稳定标准：" + standard + "</p>"
    body += "<p>仅保留白名单控件文字的脱敏弹窗截图可进入报告；原始截图、trace 和页面快照不进入报告或脚本包。本轮已清除快照文件 " + str(removed) + " 个。</p>"
    body += "<p>执行覆盖率：{:.1%}；已执行用例通过率：{:.1%}；全部用例稳定率：{:.1%}。</p>".format(s["execution_coverage"], s["pass_rate"], s["stability_rate"])
    body += '<p><a href="results.json">JSON</a> · <a href="results.xlsx">Excel</a> · <a href="scripts.zip">脚本工程</a> · <a href="asset-review.json">资产回写候选</a></p>'
    body += '<label>筛选 <input id="filter" placeholder="用例 ID、标题或状态"></label><section><table><thead><tr>'
    body += "".join("<th>" + h + "</th>" for h in ("用例ID", "来源", "标题", "进度", "结果", "稳定性", "连续成功", "脚本", "证据", "原因"))
    body += "</tr></thead><tbody>" + "".join(rows) + "</tbody></table></section><h2>耗时与人工问题</h2><pre>"
    previews = sorted((base / "batches").glob("*/preview.json"), key=lambda item: item.stat().st_mtime)
    review = base / "state/heal-review.json"
    memory = base / "state/knowledge/llms-txt.json"
    extra = {}
    if previews:
        extra["preview"] = read_json(previews[-1])
    if review.is_file():
        extra["heal_review"] = read_json(review)
    if memory.is_file():
        remembered = read_json(memory)
        extra["locator_memory"] = {"environment": remembered.get("environment"), "updated_at": remembered.get("updated_at"),
                                   "written_back": [key for key, entry in remembered.get("entries", {}).items() if entry.get("last_success")]}
    body += escape(json.dumps({"metrics": metrics, "questions": questions, "snapshots_removed": removed, **extra}, ensure_ascii=False, indent=2)) + "</pre>"
    body += '<script>document.getElementById("filter").addEventListener("input",e=>document.querySelectorAll("tbody tr").forEach(r=>r.hidden=!r.textContent.toLowerCase().includes(e.target.value.toLowerCase())))</script></html>'
    (out / "index.html").write_text(scrub_text(body, m["config"]), encoding="utf-8")
    runner = base / "runner"
    script_package = None
    if runner.exists():
        with zipfile.ZipFile(out / "scripts.zip", "w", compression=zipfile.ZIP_DEFLATED) as archive:
            from .bindings import case_bindings
            archive.writestr("runner/cases.json", scrub_text(json.dumps(case_bindings(base, m), ensure_ascii=False), m["config"]))
            protected = {Path(m["config_path"]).resolve()}
            if m["config"].get("auth", {}).get("storage_state"):
                protected.add(Path(m["config"]["auth"]["storage_state"]).resolve())
            credential = m["config"].get("auth", {}).get("credentials_file")
            if credential:
                protected.add(Path(credential).resolve())
            if m["config"].get("auth", {}).get("storage_state"):
                protected.add(Path(m["config"]["auth"]["storage_state"] + ".meta.json").resolve())
            for file in runner.rglob("*"):
                rel = file.relative_to(runner)
                if not file.is_file() or file.is_symlink() or file.resolve() in protected or any(part in {"node_modules", "results", ".auth", "test-results", "playwright-report"} for part in rel.parts):
                    continue
                if file.name.startswith(".env") or file.name in {"cases.json", "auth.json", "storage-state.json", "error-context.md", "trace.zip", "playwright-results.json"} or ".local." in file.name:
                    continue
                if file.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp", ".webm"} or any(part in {"playwright-report", "traces", "test-results"} for part in rel.parts):
                    continue
                if file.suffix in {".ts", ".js", ".mjs", ".json", ".md", ".pac", ".txt", ".yaml", ".yml"}:
                    overlay = base / "state/knowledge/llms-txt.json" if str(rel) == "knowledge/llms-txt.json" else None
                    content = overlay.read_text(encoding="utf-8") if overlay and overlay.is_file() else file.read_text(encoding="utf-8")
                    archive.writestr("runner/" + str(rel), scrub_text(content, m["config"]))
        script_package = str((out / "scripts.zip").relative_to(base))
    m["delivery"] = {"json": str((out / "results.json").relative_to(base)),
                     "html": str((out / "index.html").relative_to(base)),
                     "xlsx": str((out / "results.xlsx").relative_to(base)),
                     "asset_diff": str((out / "asset-review.json").relative_to(base)),
                     "scripts": script_package, "generated_at": now()}
    payload["delivery_complete"] = (bool(records) and all(c["work_status"] == "finished" for c in m["cases"].values())
                                    and not payload["unresolved_sheets"] and bool(script_package)
                                    and not payload["source_update_required"]
                                    and not any(q["status"] == "open" for q in m["questions"].values())
                                    and all(c.get("feedback") for c in m["cases"].values() if c.get("result") == "passed"))
    payload = json.loads(scrub_text(json.dumps(payload, ensure_ascii=False), m["config"]))
    atomic_write(out / "results.json", payload)
    return {"artifacts": m["delivery"], "summary": s, "delivery_complete": payload["delivery_complete"]}
