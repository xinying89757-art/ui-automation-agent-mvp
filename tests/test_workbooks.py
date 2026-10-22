import json
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from uia_core.workbooks import extract
from uia_core.workflow import heal_review
from uia_core import specs, runtime, common

ROOT = Path(__file__).resolve().parents[1]


class WorkbookAliasTest(unittest.TestCase):
    def test_case_description_and_level_map_to_title_and_priority(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "cases.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "测试用例"
            sheet.append(["用例编号", "模块", "用例描述", "用例级别", "前置条件", "测试步骤", "预期结果"])
            sheet.append(["TC-0001", "入口", "打开管理txt文件", "P0", "已登录", "点击设置", "看到入口"])
            book.save(path)
            parsed = extract(path)
        self.assertEqual(parsed["gaps"], [])
        case = parsed["cases"][0]
        self.assertEqual(case["title"], "打开管理txt文件")
        self.assertEqual(case["priority"], "P0")
        self.assertEqual(case["module"], "入口")
        self.assertEqual(case["suggested_id"], "TC-0001")

    def test_content_free_rows_break_continuations_and_are_reported(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "arbitrary.xlsx"
            book = Workbook(); sheet = book.active; sheet.title = "智能表1"
            sheet.append(["ID", "功能模块", "用例标题", "优先级", "前提条件", "步骤", "期望结果"])
            sheet.append(["A-1", "甲", "标题", "P1", "登录", "动作1", "结果1"])
            sheet.append([None, "乙", None, None, None, None, None])
            sheet.append([None, "甲", None, None, None, "动作2", "结果2"])
            book.save(path)
            result = extract(path)
        self.assertEqual(len(result["cases"]), 2)
        self.assertEqual(result["cases"][0]["steps"], ["动作1"])
        self.assertEqual(result["cases"][1]["steps"], ["动作2"])
        self.assertEqual(result["sheets"][0]["skipped_rows"], [{"row": 3, "reason": "no_case_content"}])

    def test_auto_id_is_stable_and_disambiguates_duplicate_titles(self):
        def parse(rows, sheet_name):
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "input.xlsx"; book = Workbook(); sheet = book.active; sheet.title = sheet_name
                sheet.append(["模块", "用例名称", "操作步骤", "预期结果"])
                for row in rows: sheet.append(row)
                book.save(path); return extract(path)["cases"]
        first = parse([[" 模块A ", "标题", "步骤一", "结果"], ["模块A", "标题", "步骤二", "结果二"]], "sheet-a")
        reordered = parse([["模块A", "标题", "步骤二", "结果二"], [" 模块A ", "标题", "步骤一", "结果"]], "renamed")
        self.assertEqual({c["suggested_id"] for c in first}, {c["suggested_id"] for c in reordered})
        self.assertEqual(first[0]["identity_key"].split(":")[0], reordered[1]["identity_key"].split(":")[0])
        unique_before = parse([["模块A", "唯一标题", "步骤", "结果"]], "x")[0]
        unique_after = parse([["模块A", "唯一标题", "步骤改", "结果"]], "x")[0]
        self.assertNotEqual(unique_before["content_hash"], unique_after["content_hash"])
        self.assertEqual(unique_before["suggested_id"], unique_after["suggested_id"])

    def test_non_ascii_id_generates_safe_spec_filename_and_grep_tag(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"; book = Workbook(); sheet = book.active
            sheet.append(["用例编号", "用例描述", "测试步骤", "预期结果"])
            sheet.append(["用例-0001", "打开窗口", "打开设置", "窗口显示"]); book.save(path)
            cid = extract(path)["cases"][0]["suggested_id"]
            self.assertEqual(cid, "C-0001")
            base = Path(directory) / "run"; (base / "cases" / cid).mkdir(parents=True); (base / "runner/tests/generated").mkdir(parents=True)
            spec = dict(specs_case := {
                "schema_version":"case-spec.v1", "case_id":cid, "intent":"打开窗口",
                "automatable":{"status":"yes","reason":"已知"}, "required_env":[], "data_prep":{},
                "mutation_scope":[], "start_page":"editor", "steps":[{"id":"S1","verb":"open","target":"settings"}],
                "assertions":[{"id":"A1","after":"S1","excel":"窗口显示","checks":[{"target":"txt.entry","op":"visible"}]}], "questions":[]})
            common.atomic_write(base / "cases" / cid / "source.json", {"case_id":cid,"expected":["窗口显示"]})
            common.atomic_write(base / "cases" / cid / "case-spec.yaml", spec, as_yaml=True)
            case = {"case_id":cid,"source":f"cases/{cid}/source.json","case_spec":f"cases/{cid}/case-spec.yaml","source_hash":"hash","attempts":[]}
            with patch.object(specs, "register_script", return_value={"script":"runner/tests/generated/C-0001.spec.ts"}):
                specs.generate(base, {"cases":{cid:case},"questions":{},"history":[]}, cid)
            self.assertTrue((base / "runner/tests/generated/C-0001.spec.ts").is_file())
            self.assertEqual(runtime.grep_tag_pattern(cid), r"\[C-0001\]")


class HealReviewTest(unittest.TestCase):
    def test_review_rejects_assertion_edit_without_rewriting_the_spec(self):
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            spec = base / "runner/tests/sample.spec.ts"
            spec.parent.mkdir(parents=True)
            original = "await expect(dialog).toContainText('故意写错的标题');\n"
            spec.write_text(original, encoding="utf-8")
            proposal = base / "evidence/proposal.ts"
            proposal.parent.mkdir(parents=True)
            proposal.write_text(original.replace("故意写错的标题", "管理txt文件"), encoding="utf-8")
            evidence = base / "evidence/review.json"
            evidence.write_text(json.dumps({
                "message": "UIA_BUSINESS_FAILURE [rehearse-title]: injected expectation",
                "spec": "runner/tests/sample.spec.ts",
                "proposal": "evidence/proposal.ts",
            }, ensure_ascii=False), encoding="utf-8")
            result = heal_review(base, "evidence/review.json")
            kept = spec.read_text(encoding="utf-8")
        self.assertFalse(result["accept"])
        self.assertEqual(result["kind"], "suspected_product_bug")
        self.assertTrue(result["spec_unchanged"])
        self.assertIn("故意写错的标题", result["rejected_diff"])
        self.assertEqual(kept, original)
        self.assertEqual(subprocess.call(["node", "--check", str(ROOT / "assets/playwright-runner/support/heal.mjs")]), 0)
